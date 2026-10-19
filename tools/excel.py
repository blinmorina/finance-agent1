"""Excel I/O tools for the finance agent.

Uses openpyxl so the agent can read/write real .xlsx workbooks with multiple
sheets, formulas, number formats, and basic styling. All paths are resolved
relative to the project root's `outputs/` folder unless an absolute path is
given.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

from agents import function_tool
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUTS_DIR = PROJECT_ROOT / "outputs"
OUTPUTS_DIR.mkdir(exist_ok=True)


def _resolve_path(path: str) -> Path:
    """Resolve a path; relative paths land in outputs/."""
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = OUTPUTS_DIR / p
    if p.suffix == "":
        p = p.with_suffix(".xlsx")
    return p


# ---------------------------------------------------------------------------
# Styling helpers (shared with dcf.py)
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
INPUT_FONT = Font(name="Calibri", size=11, color="0070C0", italic=False)
THIN = Side(border_style="thin", color="B7B7B7")
ALL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_USD = '_($* #,##0_);_($* (#,##0);_($* "-"_);_(@_)'
FMT_USD_M = '_($* #,##0.0,,_);_($* (#,##0.0,,);_($* "-"_);_(@_)'  # in millions
FMT_PCT = "0.0%"
FMT_MULT = "0.00x"
FMT_INT = "#,##0"


def style_header_row(ws, row: int, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autosize_columns(ws, min_width: int = 12, max_width: int = 40) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = min_width
        for cell in col:
            try:
                v = cell.value
                if v is None:
                    continue
                s = str(v)
                if len(s) > max_len:
                    max_len = len(s)
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(max_len + 2, max_width)


# ---------------------------------------------------------------------------
# Agent tools
# ---------------------------------------------------------------------------


@function_tool
def list_excel_files() -> str:
    """List all Excel files (.xlsx) available in the outputs/ directory.

    Returns a JSON string with filenames and sizes. Use this to discover what
    workbooks already exist before reading or modifying them.
    """
    files = []
    for p in sorted(OUTPUTS_DIR.glob("*.xlsx")):
        files.append({"name": p.name, "size_bytes": p.stat().st_size, "path": str(p)})
    return json.dumps({"outputs_dir": str(OUTPUTS_DIR), "files": files}, indent=2)


@function_tool
def read_excel(file_path: str, sheet_name: str | None = None, max_rows: int = 200) -> str:
    """Read an Excel workbook and return its contents as JSON.

    Args:
        file_path: Path to the .xlsx file. Relative paths resolve into outputs/.
        sheet_name: Optional specific sheet to read. If omitted, returns all sheets.
        max_rows: Maximum rows per sheet to return (to keep output manageable).

    Returns JSON with sheet names, dimensions, and cell values. Formulas are
    returned as their stored formula string (e.g. "=SUM(A1:A5)").
    """
    path = _resolve_path(file_path)
    if not path.exists():
        return json.dumps({"error": f"File not found: {path}"})

    wb = load_workbook(path, data_only=False)
    result: dict[str, Any] = {"path": str(path), "sheets": {}}

    sheets = [sheet_name] if sheet_name else wb.sheetnames
    for name in sheets:
        if name not in wb.sheetnames:
            result["sheets"][name] = {"error": f"sheet not found. available: {wb.sheetnames}"}
            continue
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                rows.append([f"... ({ws.max_row - max_rows} more rows truncated)"])
                break
            rows.append(list(row))
        result["sheets"][name] = {
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "rows": rows,
        }
    return json.dumps(result, default=str, indent=2)


@function_tool(strict_mode=False)
def write_excel_sheet(
    file_path: str,
    sheet_name: str,
    rows: list[list[Any]],
    overwrite_sheet: bool = True,
    header_row: bool = True,
) -> str:
    """Write a 2D array of values/formulas to a sheet in an Excel workbook.

    Cells whose string values start with '=' are written as formulas (e.g.
    "=SUM(B2:B10)"). The workbook is created if it doesn't exist. If the
    workbook exists, this updates/adds the named sheet.

    Args:
        file_path: Path to the .xlsx file. Relative paths resolve into outputs/.
        sheet_name: Name of the sheet to write.
        rows: 2D list of cell values. Strings starting with "=" become formulas.
        overwrite_sheet: If True and the sheet already exists, replace it.
        header_row: If True, the first row is styled as a header.

    Returns a JSON status with the absolute path written.
    """
    path = _resolve_path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        wb = load_workbook(path)
        if sheet_name in wb.sheetnames:
            if overwrite_sheet:
                del wb[sheet_name]
            else:
                return json.dumps(
                    {"error": f"sheet '{sheet_name}' exists; pass overwrite_sheet=true to replace"}
                )
        ws = wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        ws = wb.create_sheet(sheet_name)

    last_col = 0
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
            if c > last_col:
                last_col = c

    if header_row and rows:
        style_header_row(ws, 1, last_col)
        ws.freeze_panes = "A2"

    autosize_columns(ws)
    wb.save(path)
    return json.dumps(
        {
            "status": "ok",
            "path": str(path),
            "sheet": sheet_name,
            "rows_written": len(rows),
        }
    )


@function_tool(strict_mode=False)
def update_excel_cells(file_path: str, sheet_name: str, updates: dict[str, Any]) -> str:
    """Update specific cells in a sheet by A1 notation.

    Args:
        file_path: Path to the .xlsx file.
        sheet_name: The sheet to modify.
        updates: Mapping of A1 addresses to values, e.g. {"B2": 100, "C5": "=B2*1.1"}.
            Values starting with "=" are treated as formulas.

    Returns JSON status.
    """
    path = _resolve_path(file_path)
    if not path.exists():
        return json.dumps({"error": f"file not found: {path}"})
    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        return json.dumps({"error": f"sheet '{sheet_name}' not in {wb.sheetnames}"})
    ws = wb[sheet_name]
    for addr, val in updates.items():
        ws[addr] = val
    wb.save(path)
    return json.dumps({"status": "ok", "path": str(path), "updated": list(updates.keys())})


@function_tool
def add_named_range(file_path: str, name: str, sheet_name: str, cell_range: str) -> str:
    """Create a workbook-level named range (e.g. WACC -> Inputs!$B$3).

    Useful in financial models so formulas can reference inputs by name.
    """
    from openpyxl.workbook.defined_name import DefinedName

    path = _resolve_path(file_path)
    if not path.exists():
        return json.dumps({"error": f"file not found: {path}"})
    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        return json.dumps({"error": f"sheet '{sheet_name}' not in {wb.sheetnames}"})

    ref = f"'{sheet_name}'!{cell_range}"
    dn = DefinedName(name=name, attr_text=ref)
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names[name] = dn
    wb.save(path)
    return json.dumps({"status": "ok", "name": name, "refers_to": ref})
