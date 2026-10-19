"""DCF (Discounted Cash Flow) model builder.

Produces a fully-formed Excel workbook with linked formulas across:
  - Cover: summary outputs
  - Assumptions: operating inputs, terminal method, growth path (5–10 years)
  - WACC_Build: CAPM + Hamada unlevering/relevering → WACC (optional manual override)
  - DCF: revenue → FCF → PV; Gordon growth / exit-multiple / blend terminal EV
  - Sensitivity: static WACC × terminal growth grid (Gordon path)

Everything projects as Excel formulas except the sensitivity snapshot cells.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import yfinance as yf
from agents import function_tool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .excel import (
    ALL_BORDER,
    BOLD_FONT,
    FMT_MULT,
    FMT_PCT,
    FMT_USD,
    FMT_USD_M,
    HEADER_FILL,
    HEADER_FONT,
    INPUT_FILL,
    INPUT_FONT,
    OUTPUTS_DIR,
    SUBHEADER_FILL,
    TOTAL_FILL,
    autosize_columns,
)


def _input_cell(ws, addr: str, value: Any, fmt: str | None = None) -> None:
    """Mark a cell as a user-editable input (yellow fill, blue font)."""
    cell = ws[addr]
    cell.value = value
    cell.fill = INPUT_FILL
    cell.font = INPUT_FONT
    cell.border = ALL_BORDER
    if fmt:
        cell.number_format = fmt


def _label(ws, addr: str, text: str, bold: bool = False) -> None:
    cell = ws[addr]
    cell.value = text
    if bold:
        cell.font = BOLD_FONT


def _formula(ws, addr: str, formula: str, fmt: str | None = None, bold: bool = False) -> None:
    cell = ws[addr]
    cell.value = formula
    if fmt:
        cell.number_format = fmt
    if bold:
        cell.font = BOLD_FONT


def _add_name(wb: Workbook, name: str, ref: str) -> None:
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names[name] = DefinedName(name=name, attr_text=ref)


def _try_pull_inputs(ticker: str) -> dict[str, Any]:
    """Pull a sensible default set of inputs from yfinance. All optional."""
    try:
        t = yf.Ticker(ticker.upper())
        info = t.info or {}
        inc = t.income_stmt
        bs = t.balance_sheet

        def _row(df, names: list[str]):
            if df is None or df.empty:
                return None
            for n in names:
                if n in df.index:
                    try:
                        v = df.iloc[df.index.get_loc(n), 0]
                        if v == v:  # NaN check
                            return float(v)
                    except Exception:
                        pass
            return None

        revenue = _row(inc, ["Total Revenue", "Revenue", "Operating Revenue"])
        op_income = _row(inc, ["Operating Income", "Operating Income Loss"])
        debt = _row(bs, ["Total Debt"]) or 0.0
        cash = _row(bs, ["Cash And Cash Equivalents", "Cash Cash Equivalents And Short Term Investments"]) or 0.0

        ebit_margin = (op_income / revenue) if (op_income and revenue) else 0.20
        shares = info.get("sharesOutstanding") or 0
        beta = info.get("beta") or 1.0
        rev_growth = info.get("revenueGrowth") or 0.08
        price = info.get("currentPrice") or info.get("regularMarketPrice")
        mcap = info.get("marketCap")
        try:
            if mcap is None and shares and price:
                mcap = float(shares) * float(price)
        except Exception:
            mcap = None
        nd = float(debt - cash)
        debt_weight = None
        if mcap and float(mcap) > 0:
            ev_approx = nd + float(mcap)
            if ev_approx > 0:
                debt_weight = max(0.05, min(0.80, nd / ev_approx))

        return {
            "company_name": info.get("longName") or ticker.upper(),
            "currency": info.get("currency") or "USD",
            "base_revenue": revenue or 1_000_000_000,
            "ebit_margin": max(0.05, min(0.45, ebit_margin)),
            "revenue_growth": max(0.0, min(0.30, rev_growth)),
            "net_debt": nd,
            "shares_outstanding": shares,
            "beta": beta,
            "current_price": price,
            "market_cap": float(mcap) if mcap is not None else None,
            "debt_weight_wacc": float(debt_weight) if debt_weight is not None else 0.25,
        }
    except Exception:
        return {
            "company_name": ticker.upper(),
            "currency": "USD",
            "base_revenue": 1_000_000_000,
            "ebit_margin": 0.20,
            "revenue_growth": 0.08,
            "net_debt": 0.0,
            "shares_outstanding": 1_000_000,
            "beta": 1.0,
            "current_price": None,
            "market_cap": None,
            "debt_weight_wacc": 0.25,
        }


def _computed_wacc(
    rf: float,
    erp: float,
    levered_beta: float,
    wd_wacc: float,
    wd_beta: float,
    kd: float,
    tax: float,
) -> float:
    """Hamada unlever / relever at target capital structure; classic WACC."""
    wd_wacc = max(1e-6, min(0.95, wd_wacc))
    wd_beta = max(1e-6, min(0.95, wd_beta))
    de_beta = wd_beta / (1 - wd_beta)
    bu = levered_beta / (1 + (1 - tax) * de_beta)
    de_wacc = wd_wacc / (1 - wd_wacc)
    bl = bu * (1 + (1 - tax) * de_wacc)
    ke = rf + bl * erp
    return (1 - wd_wacc) * ke + wd_wacc * kd * (1 - tax)


@function_tool(strict_mode=False)
def build_dcf_model(
    ticker: str,
    output_file: str | None = None,
    base_revenue: float | None = None,
    revenue_growth_rates: list[float] | None = None,
    ebit_margin: float | None = None,
    tax_rate: float = 0.21,
    da_pct_revenue: float = 0.04,
    capex_pct_revenue: float = 0.05,
    nwc_pct_revenue: float = 0.02,
    wacc: float | None = None,
    terminal_growth: float = 0.025,
    terminal_exit_ebitda_multiple: float | None = None,
    terminal_value_method: int = 1,
    risk_free_rate: float = 0.042,
    equity_risk_premium: float = 0.055,
    levered_beta: float | None = None,
    debt_weight_wacc: float | None = None,
    debt_weight_for_beta: float | None = None,
    cost_of_debt_pretax: float = 0.065,
    net_debt: float | None = None,
    shares_outstanding: float | None = None,
    forecast_years: int = 5,
) -> str:
    """Build a DCF workbook (5–10 forecast years) with WACC build + dual terminal methods.

    Any inputs left as None will be auto-populated from market data (yfinance)
    when possible.

    Terminal methods (Assumptions ``terminal_value_method``): 1 = Gordon growth only,
    2 = exit enterprise value on EBITDA × multiple only, 3 = simple average of both.

    ``wacc``: when provided, overrides the WACC_Build formula output (manual discount rate).
    """
    defaults = _try_pull_inputs(ticker)
    fy = min(max(int(forecast_years), 5), 10)

    base_revenue = base_revenue if base_revenue is not None else defaults["base_revenue"]
    ebit_margin = ebit_margin if ebit_margin is not None else defaults["ebit_margin"]
    net_debt = net_debt if net_debt is not None else defaults["net_debt"]
    shares_outstanding = (
        shares_outstanding if shares_outstanding is not None else defaults["shares_outstanding"]
    )
    bl = float(levered_beta) if levered_beta is not None else float(defaults["beta"])
    wd_w = float(debt_weight_wacc) if debt_weight_wacc is not None else float(defaults["debt_weight_wacc"])
    wd_b = float(debt_weight_for_beta) if debt_weight_for_beta is not None else wd_w
    term_mult = float(terminal_exit_ebitda_multiple) if terminal_exit_ebitda_multiple is not None else 11.0
    tvm = min(max(int(terminal_value_method), 1), 3)

    if revenue_growth_rates is None:
        g = defaults["revenue_growth"]
        revenue_growth_rates = [g, g * 0.85, g * 0.7, g * 0.55, g * 0.4, g * 0.35, g * 0.3, g * 0.28, g * 0.26, g * 0.24][:fy]
    elif isinstance(revenue_growth_rates, (int, float)):
        revenue_growth_rates = [float(revenue_growth_rates)] * fy
    if len(revenue_growth_rates) < fy:
        last = revenue_growth_rates[-1]
        revenue_growth_rates += [last] * (fy - len(revenue_growth_rates))
    revenue_growth_rates = revenue_growth_rates[:fy]

    fname = output_file or f"{ticker.upper()}_DCF.xlsx"
    if not fname.endswith(".xlsx"):
        fname += ".xlsx"
    path = Path(fname)
    if not path.is_absolute():
        path = OUTPUTS_DIR / path

    tax_f = float(tax_rate)
    computed_w = _computed_wacc(
        float(risk_free_rate),
        float(equity_risk_premium),
        bl,
        wd_w,
        wd_b,
        float(cost_of_debt_pretax),
        tax_f,
    )
    final_wacc = float(wacc) if wacc is not None else computed_w

    wb = Workbook()
    wb.remove(wb.active)

    max_col_ltr = get_column_letter(2 + fy)
    merge_w = get_column_letter(max(7, 2 + fy))

    # ------------------------------------------------------------------
    # ASSUMPTIONS SHEET
    # ------------------------------------------------------------------
    a = wb.create_sheet("Assumptions")
    a["A1"] = f"{defaults['company_name']} — DCF Assumptions"
    a["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    a.merge_cells(f"A1:{merge_w}1")

    a["A3"] = "Operating inputs (yellow where marked)"
    a["A3"].font = BOLD_FONT
    a["A3"].fill = SUBHEADER_FILL
    a.merge_cells(f"A3:{merge_w}3")

    _label(a, "A4", "Ticker", bold=True)
    a["B4"] = ticker.upper()
    _label(a, "A5", "Currency", bold=True)
    a["B5"] = defaults["currency"]

    _label(a, "A7", "Base year revenue", bold=True)
    _input_cell(a, "B7", float(base_revenue), FMT_USD)

    _label(a, "A8", "Tax rate", bold=True)
    _input_cell(a, "B8", float(tax_rate), FMT_PCT)

    _label(a, "A9", "EBIT margin (steady)", bold=True)
    _input_cell(a, "B9", float(ebit_margin), FMT_PCT)

    _label(a, "A10", "D&A % of revenue", bold=True)
    _input_cell(a, "B10", float(da_pct_revenue), FMT_PCT)

    _label(a, "A11", "CapEx % of revenue", bold=True)
    _input_cell(a, "B11", float(capex_pct_revenue), FMT_PCT)

    _label(a, "A12", "Δ NWC % of revenue", bold=True)
    _input_cell(a, "B12", float(nwc_pct_revenue), FMT_PCT)

    _label(a, "A13", "WACC (linked to WACC_Build)", bold=True)
    _formula(a, "B13", "=WACC_Build!$B$21", FMT_PCT, bold=True)

    _label(a, "A14", "Terminal growth rate (Gordon)", bold=True)
    _input_cell(a, "B14", float(terminal_growth), FMT_PCT)

    _label(a, "A15", "Terminal exit EV / EBITDA multiple", bold=True)
    _input_cell(a, "B15", float(term_mult), FMT_MULT)

    _label(a, "A16", "Terminal method (1=Gordon, 2=Multiple, 3=blend avg)", bold=True)
    _input_cell(a, "B16", int(tvm), "0")

    _label(a, "A17", "Net debt", bold=True)
    _input_cell(a, "B17", float(net_debt), FMT_USD)

    _label(a, "A18", "Shares outstanding", bold=True)
    _input_cell(a, "B18", float(shares_outstanding), "#,##0")

    gr_hdr_row = 21
    gr_period_row = 22
    gr_growth_row = 23
    _label(a, f"A{gr_hdr_row}", "Revenue growth by year", bold=True)
    a[f"A{gr_hdr_row}"].fill = SUBHEADER_FILL
    a.merge_cells(f"A{gr_hdr_row}:{max_col_ltr}{gr_hdr_row}")
    for i, g in enumerate(revenue_growth_rates):
        col = get_column_letter(2 + i)
        a[f"{col}{gr_period_row}"] = f"Year {i+1}"
        a[f"{col}{gr_period_row}"].font = BOLD_FONT
        a[f"{col}{gr_period_row}"].alignment = Alignment(horizontal="center")
        _input_cell(a, f"{col}{gr_growth_row}", float(g), FMT_PCT)
    a[f"A{gr_period_row}"] = "Period"
    a[f"A{gr_period_row}"].font = BOLD_FONT
    a[f"A{gr_growth_row}"] = "Growth"
    a[f"A{gr_growth_row}"].font = BOLD_FONT

    # ------------------------------------------------------------------
    # WACC BUILD
    # ------------------------------------------------------------------
    wsheet = wb.create_sheet("WACC_Build")
    wsheet["A1"] = "WACC — CAPM + Hamada unlevering / relevering"
    wsheet["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    wsheet.merge_cells("A1:E1")

    wsheet["A3"] = "Inputs"
    wsheet["A3"].font = BOLD_FONT
    wsheet["A3"].fill = SUBHEADER_FILL

    _label(wsheet, "A4", "Risk-free rate", bold=True)
    _input_cell(wsheet, "B4", float(risk_free_rate), FMT_PCT)
    _label(wsheet, "A5", "Equity risk premium", bold=True)
    _input_cell(wsheet, "B5", float(equity_risk_premium), FMT_PCT)
    _label(wsheet, "A6", "Levered equity beta (observed)", bold=True)
    _input_cell(wsheet, "B6", float(bl), FMT_MULT)
    _label(wsheet, "A7", "Debt weight for WACC (D / V)", bold=True)
    _input_cell(wsheet, "B7", float(wd_w), FMT_PCT)
    _label(wsheet, "A8", "Debt weight for unlevering beta (D / V)", bold=True)
    _input_cell(wsheet, "B8", float(wd_b), FMT_PCT)
    _label(wsheet, "A9", "Pre-tax cost of debt", bold=True)
    _input_cell(wsheet, "B9", float(cost_of_debt_pretax), FMT_PCT)
    _label(wsheet, "A10", "Tax rate (= Assumptions)", bold=True)
    _formula(wsheet, "B10", "=Assumptions!$B$8", FMT_PCT, bold=True)

    _label(wsheet, "A12", "D/E for beta observation", bold=True)
    _formula(wsheet, "B12", "=B8/(1-B8)", FMT_MULT)
    _label(wsheet, "A13", "D/E for target capital structure", bold=True)
    _formula(wsheet, "B13", "=B7/(1-B7)", FMT_MULT)
    _label(wsheet, "A14", "Unlevered beta (Bu)", bold=True)
    _formula(wsheet, "B14", "=B6/(1+(1-B10)*B12)", FMT_MULT, bold=True)
    _label(wsheet, "A15", "Relevered beta at target (Bl)", bold=True)
    _formula(wsheet, "B15", "=B14*(1+(1-B10)*B13)", FMT_MULT, bold=True)
    _label(wsheet, "A16", "Cost of equity Ke = Rf + Bl × ERP", bold=True)
    _formula(wsheet, "B16", "=B4+B15*B5", FMT_PCT, bold=True)
    _label(wsheet, "A17", "WACC (formula)", bold=True)
    _formula(wsheet, "B17", "=(1-B7)*B16+B7*B9*(1-B10)", FMT_PCT, bold=True)

    _label(wsheet, "A19", "Manual WACC override (0 = use formula)", bold=True)
    _input_cell(wsheet, "B19", float(wacc) if wacc is not None else 0.0, FMT_PCT)
    _label(wsheet, "A20", "Selected discount rate", bold=True)
    _formula(wsheet, "B20", "=IF(B19>0,B19,B17)", FMT_PCT, bold=True)
    _label(wsheet, "A21", "Final WACC (→ Assumptions!B13)", bold=True)
    _formula(wsheet, "B21", "=B20", FMT_PCT, bold=True)

    autosize_columns(wsheet)

    # Named ranges for DCF formulas
    _add_name(wb, "BaseRevenue", "Assumptions!$B$7")
    _add_name(wb, "TaxRate", "Assumptions!$B$8")
    _add_name(wb, "EBITMargin", "Assumptions!$B$9")
    _add_name(wb, "DA_Pct", "Assumptions!$B$10")
    _add_name(wb, "Capex_Pct", "Assumptions!$B$11")
    _add_name(wb, "NWC_Pct", "Assumptions!$B$12")
    _add_name(wb, "WACC", "Assumptions!$B$13")
    _add_name(wb, "TGrowth", "Assumptions!$B$14")
    _add_name(wb, "ExitMult", "Assumptions!$B$15")
    _add_name(wb, "TermMethod", "Assumptions!$B$16")
    _add_name(wb, "NetDebt", "Assumptions!$B$17")
    _add_name(wb, "Shares", "Assumptions!$B$18")

    autosize_columns(a)

    # ------------------------------------------------------------------
    # DCF SHEET — projection model
    # ------------------------------------------------------------------
    d = wb.create_sheet("DCF")
    d["A1"] = f"{defaults['company_name']} — {fy}-Year DCF Model"
    d["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    d.merge_cells(f"A1:{get_column_letter(2 + fy)}1")

    d["A3"] = "Line Item"
    d["A3"].font = HEADER_FONT
    d["A3"].fill = HEADER_FILL
    for i in range(fy):
        col = get_column_letter(2 + i)
        d[f"{col}3"] = f"Year {i+1}"
        d[f"{col}3"].font = HEADER_FONT
        d[f"{col}3"].fill = HEADER_FILL
        d[f"{col}3"].alignment = Alignment(horizontal="center")
    term_col = get_column_letter(2 + fy)
    d[f"{term_col}3"] = "Terminal"
    d[f"{term_col}3"].font = HEADER_FONT
    d[f"{term_col}3"].fill = HEADER_FILL
    d[f"{term_col}3"].alignment = Alignment(horizontal="center")

    gr_row = gr_growth_row
    _label(d, "A4", "Revenue", bold=True)
    for i in range(fy):
        col = get_column_letter(2 + i)
        growth_col = get_column_letter(2 + i)
        if i == 0:
            _formula(d, f"{col}4", f"=BaseRevenue*(1+Assumptions!{growth_col}{gr_row})", FMT_USD_M)
        else:
            prev = get_column_letter(2 + i - 1)
            _formula(d, f"{col}4", f"={prev}4*(1+Assumptions!{growth_col}{gr_row})", FMT_USD_M)
    last_fc_col = get_column_letter(2 + fy - 1)
    _formula(d, f"{term_col}4", f"={last_fc_col}4*(1+TGrowth)", FMT_USD_M)

    _label(d, "A5", "EBIT", bold=True)
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}5", f"={col}4*EBITMargin", FMT_USD_M)

    _label(d, "A6", "Taxes")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}6", f"=-{col}5*TaxRate", FMT_USD_M)

    _label(d, "A7", "NOPAT", bold=True)
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}7", f"={col}5+{col}6", FMT_USD_M, bold=True)

    _label(d, "A8", "+ D&A")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}8", f"={col}4*DA_Pct", FMT_USD_M)

    _label(d, "A9", "EBITDA (for exit multiple)")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}9", f"={col}5+{col}8", FMT_USD_M)

    _label(d, "A10", "− CapEx")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}10", f"=-{col}4*Capex_Pct", FMT_USD_M)

    _label(d, "A11", "− Δ NWC")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        if i == 0:
            _formula(d, f"{col}11", f"=-({col}4-BaseRevenue)*NWC_Pct", FMT_USD_M)
        else:
            prev = get_column_letter(2 + i - 1)
            _formula(d, f"{col}11", f"=-({col}4-{prev}4)*NWC_Pct", FMT_USD_M)

    _label(d, "A12", "Free Cash Flow", bold=True)
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        cell = d[f"{col}12"]
        cell.value = f"={col}7+{col}8+{col}10+{col}11"
        cell.number_format = FMT_USD_M
        cell.font = BOLD_FONT
        cell.fill = TOTAL_FILL

    _label(d, "A13", "Discount period")
    for i in range(fy):
        col = get_column_letter(2 + i)
        d[f"{col}13"] = i + 1
        d[f"{col}13"].alignment = Alignment(horizontal="center")
    d[f"{term_col}13"] = fy
    d[f"{term_col}13"].alignment = Alignment(horizontal="center")

    _label(d, "A14", "Discount factor")
    for i in range(fy + 1):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}14", f"=1/(1+WACC)^{col}13", "0.0000")

    _label(d, "A15", "Terminal EV — Gordon growth")
    _formula(
        d,
        f"{term_col}15",
        f"=IF(WACC<=TGrowth,0,{term_col}12/(WACC-TGrowth))",
        FMT_USD_M,
        bold=True,
    )

    _label(d, "A16", "Terminal EV — exit multiple × EBITDA")
    _formula(d, f"{term_col}16", f"={term_col}9*ExitMult", FMT_USD_M, bold=True)

    _label(d, "A17", "Terminal EV — selected")
    _formula(
        d,
        f"{term_col}17",
        f"=CHOOSE(TermMethod,{term_col}15,{term_col}16,({term_col}15+{term_col}16)/2)",
        FMT_USD_M,
        bold=True,
    )

    _label(d, "A18", "Present value")
    for i in range(fy):
        col = get_column_letter(2 + i)
        _formula(d, f"{col}18", f"={col}12*{col}14", FMT_USD_M, bold=True)
    _formula(d, f"{term_col}18", f"={term_col}17*{term_col}14", FMT_USD_M, bold=True)

    _label(d, "A20", "Valuation Summary", bold=True)
    d["A20"].fill = SUBHEADER_FILL
    d.merge_cells(f"A20:{term_col}20")

    fc_range = f"B18:{last_fc_col}18"
    _label(d, "A21", "Sum of PV (forecast)", bold=True)
    _formula(d, "B21", f"=SUM({fc_range})", FMT_USD_M, bold=True)
    _label(d, "A22", "PV of terminal enterprise value", bold=True)
    _formula(d, "B22", f"={term_col}18", FMT_USD_M, bold=True)
    _label(d, "A23", "Enterprise value", bold=True)
    _formula(d, "B23", "=B21+B22", FMT_USD_M, bold=True)
    d["B23"].fill = TOTAL_FILL
    _label(d, "A24", "(−) Net debt", bold=True)
    _formula(d, "B24", "=-NetDebt", FMT_USD_M)
    _label(d, "A25", "Equity value", bold=True)
    _formula(d, "B25", "=B23+B24", FMT_USD_M, bold=True)
    d["B25"].fill = TOTAL_FILL
    _label(d, "A26", "Shares outstanding", bold=True)
    _formula(d, "B26", "=Shares", "#,##0")
    _label(d, "A27", "Implied share price", bold=True)
    _formula(d, "B27", "=B25/B26", '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', bold=True)
    d["B27"].fill = TOTAL_FILL

    _add_name(wb, "EnterpriseValue", "DCF!$B$23")
    _add_name(wb, "EquityValue", "DCF!$B$25")
    _add_name(wb, "ImpliedPrice", "DCF!$B$27")

    autosize_columns(d, min_width=14)

    # ------------------------------------------------------------------
    # SENSITIVITY SHEET — WACC vs Terminal Growth (Gordon path only)
    # ------------------------------------------------------------------
    s = wb.create_sheet("Sensitivity")
    s["A1"] = "Sensitivity: Implied Share Price (Gordon terminal path)"
    s["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    s.merge_cells("A1:H1")
    s["A2"] = (
        "Rows = WACC, Cols = Terminal growth. Static snapshot — uses Gordon terminal only; "
        "re-run tool after changing exit multiples or blend mode."
    )
    s["A2"].font = Font(italic=True, color="666666")
    s.merge_cells("A2:H2")

    wacc_steps = [
        max(0.03, final_wacc - 0.02),
        max(0.03, final_wacc - 0.01),
        final_wacc,
        final_wacc + 0.01,
        final_wacc + 0.02,
    ]
    g_steps = [
        max(-0.02, terminal_growth - 0.01),
        max(-0.02, terminal_growth - 0.005),
        terminal_growth,
        terminal_growth + 0.005,
        min(0.06, terminal_growth + 0.01),
    ]

    s["A4"] = "WACC \\ g"
    s["A4"].font = HEADER_FONT
    s["A4"].fill = HEADER_FILL
    for j, gv in enumerate(g_steps):
        col = get_column_letter(2 + j)
        s[f"{col}4"] = gv
        s[f"{col}4"].number_format = FMT_PCT
        s[f"{col}4"].font = HEADER_FONT
        s[f"{col}4"].fill = HEADER_FILL
        s[f"{col}4"].alignment = Alignment(horizontal="center")

    base_rev = float(base_revenue)
    margin = float(ebit_margin)
    tax = float(tax_rate)
    da = float(da_pct_revenue)
    capex = float(capex_pct_revenue)
    nwc = float(nwc_pct_revenue)
    nd = float(net_debt)
    sh = float(shares_outstanding) or 1.0
    growths = list(revenue_growth_rates)

    for i, w in enumerate(wacc_steps):
        row = 5 + i
        s[f"A{row}"] = w
        s[f"A{row}"].number_format = FMT_PCT
        s[f"A{row}"].font = HEADER_FONT
        s[f"A{row}"].fill = HEADER_FILL
        for j, gv in enumerate(g_steps):
            col = get_column_letter(2 + j)
            price = _compute_share_price(
                base_rev,
                growths,
                margin,
                tax,
                da,
                capex,
                nwc,
                w,
                gv,
                nd,
                sh,
                terminal_exit_multiple=term_mult,
                terminal_method=1,
            )
            s[f"{col}{row}"] = price
            s[f"{col}{row}"].number_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
            s[f"{col}{row}"].border = ALL_BORDER
            if abs(w - final_wacc) < 1e-9 and abs(gv - terminal_growth) < 1e-9 and tvm == 1:
                s[f"{col}{row}"].fill = TOTAL_FILL
                s[f"{col}{row}"].font = BOLD_FONT

    autosize_columns(s, min_width=12)

    # ------------------------------------------------------------------
    # COVER SHEET
    # ------------------------------------------------------------------
    c = wb.create_sheet("Cover", 0)
    c["A1"] = f"{defaults['company_name']} ({ticker.upper()})"
    c["A1"].font = Font(name="Calibri", size=20, bold=True, color="1F4E78")
    c.merge_cells("A1:F1")
    c["A2"] = "Discounted Cash Flow Valuation Model"
    c["A2"].font = Font(name="Calibri", size=12, italic=True, color="666666")
    c.merge_cells("A2:F2")

    c["A4"] = "Key Outputs"
    c["A4"].font = BOLD_FONT
    c["A4"].fill = SUBHEADER_FILL
    c.merge_cells("A4:F4")

    _label(c, "A5", "Enterprise value", bold=True)
    _formula(c, "B5", "=EnterpriseValue", FMT_USD_M, bold=True)
    _label(c, "A6", "Equity value", bold=True)
    _formula(c, "B6", "=EquityValue", FMT_USD_M, bold=True)
    _label(c, "A7", "Implied share price", bold=True)
    _formula(c, "B7", "=ImpliedPrice", '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', bold=True)
    c["B7"].fill = TOTAL_FILL

    current_price = defaults.get("current_price")
    if current_price:
        _label(c, "A8", "Current market price", bold=True)
        c["B8"] = float(current_price)
        c["B8"].number_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
        _label(c, "A9", "Implied upside / (downside)", bold=True)
        c["B9"] = f"=B7/B8-1"
        c["B9"].number_format = FMT_PCT
        c["B9"].font = BOLD_FONT

    c["A11"] = "How to use this model"
    c["A11"].font = BOLD_FONT
    c["A11"].fill = SUBHEADER_FILL
    c.merge_cells("A11:F11")
    notes = [
        "1. Yellow inputs live on Assumptions and WACC_Build (override WACC with B19>0 on WACC_Build).",
        f"2. Forecast horizon is {fy} years; terminal column applies TGrowth to revenue then builds EBITDA-based exit multiples.",
        "3. Terminal enterprise value: Gordon FCF perpetuity, EBITDA × exit multiple, or blend — selected via Assumptions B16.",
        "4. The Sensitivity grid is a static Gordon-only snapshot — rebuild to refresh.",
        "5. Defaults come from market data; verify before relying on outputs.",
    ]
    for i, n in enumerate(notes):
        c[f"A{12+i}"] = n
        c.merge_cells(f"A{12+i}:F{12+i}")

    autosize_columns(c, min_width=22)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    implied_price = _compute_share_price(
        base_rev,
        growths,
        margin,
        tax,
        da,
        capex,
        nwc,
        final_wacc,
        float(terminal_growth),
        nd,
        sh,
        terminal_exit_multiple=term_mult,
        terminal_method=tvm,
    )

    return json.dumps(
        {
            "status": "ok",
            "path": str(path),
            "ticker": ticker.upper(),
            "summary": {
                "implied_share_price": implied_price,
                "current_price": current_price,
                "upside_pct": (implied_price / current_price - 1) * 100 if current_price else None,
                "assumptions": {
                    "forecast_years": fy,
                    "base_revenue": base_rev,
                    "revenue_growth_rates": growths,
                    "ebit_margin": margin,
                    "tax_rate": tax,
                    "wacc_computed": computed_w,
                    "wacc_final": final_wacc,
                    "wacc_manual_override": wacc,
                    "terminal_growth": float(terminal_growth),
                    "terminal_exit_ebitda_multiple": term_mult,
                    "terminal_value_method": tvm,
                    "risk_free_rate": float(risk_free_rate),
                    "equity_risk_premium": float(equity_risk_premium),
                    "levered_beta": bl,
                    "debt_weight_wacc": wd_w,
                    "debt_weight_for_beta": wd_b,
                    "cost_of_debt_pretax": float(cost_of_debt_pretax),
                    "net_debt": nd,
                    "shares_outstanding": sh,
                },
            },
            "next_steps": "Open in Excel. Flex Assumptions + WACC_Build; audit with audit_excel_model.",
        },
        default=str,
        indent=2,
    )


def _compute_share_price(
    base_rev: float,
    growths: list[float],
    margin: float,
    tax: float,
    da: float,
    capex: float,
    nwc: float,
    wacc: float,
    g_term: float,
    net_debt: float,
    shares: float,
    *,
    terminal_exit_multiple: float = 11.0,
    terminal_method: int = 1,
) -> float:
    """Plain-Python DCF mirror for sensitivity grid and JSON headline."""
    revenues: list[float] = []
    prev = base_rev
    for g in growths:
        nxt = prev * (1 + g)
        revenues.append(nxt)
        prev = nxt

    fcfs = []
    prev_rev = base_rev
    for rev in revenues:
        ebit = rev * margin
        nopat = ebit * (1 - tax)
        da_v = rev * da
        capex_v = rev * capex
        nwc_v = (rev - prev_rev) * nwc
        fcf = nopat + da_v - capex_v - nwc_v
        fcfs.append(fcf)
        prev_rev = rev

    pv_fcf = sum(f / (1 + wacc) ** (i + 1) for i, f in enumerate(fcfs))

    terminal_rev = revenues[-1] * (1 + g_term)
    terminal_ebit = terminal_rev * margin
    terminal_nopat = terminal_ebit * (1 - tax)
    terminal_fcf = (
        terminal_nopat
        + terminal_rev * da
        - terminal_rev * capex
        - (terminal_rev - revenues[-1]) * nwc
    )

    ebitda_term = terminal_ebit + terminal_rev * da
    tv_gordon = terminal_fcf / (wacc - g_term) if wacc > g_term else float("nan")
    tv_multiple = ebitda_term * terminal_exit_multiple

    if terminal_method == 1:
        tv = tv_gordon
    elif terminal_method == 2:
        tv = tv_multiple
    else:
        g_ok = tv_gordon == tv_gordon and wacc > g_term
        tv = 0.5 * (tv_gordon + tv_multiple) if g_ok else tv_multiple

    pv_tv = tv / (1 + wacc) ** len(growths)
    ev = pv_fcf + pv_tv
    equity = ev - net_debt
    return equity / shares if shares else float("nan")
