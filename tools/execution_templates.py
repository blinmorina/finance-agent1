"""PE deal execution Excel templates—precedent transactions, diligence tracking,
financing comparison, portfolio monitoring, and covenant review workbooks.

All paths default to ./outputs/. Inputs use JSON strings for flexible structures
(strict_mode=False for Agents SDK compatibility).
"""

from __future__ import annotations

import json
import statistics
from pathlib import Path
from typing import Any

from agents import function_tool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .excel import (
    ALL_BORDER,
    BOLD_FONT,
    FMT_MULT,
    FMT_PCT,
    FMT_USD,
    HEADER_FILL,
    HEADER_FONT,
    INPUT_FILL,
    INPUT_FONT,
    OUTPUTS_DIR,
    autosize_columns,
    style_header_row,
)


def _out_path(filename: str | None, default_name: str) -> Path:
    fname = filename or default_name
    if not fname.endswith(".xlsx"):
        fname += ".xlsx"
    path = Path(fname)
    if not path.is_absolute():
        path = OUTPUTS_DIR / path
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _input_cell(ws, row: int, col: int, value: Any, fmt: str | None = None) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.fill = INPUT_FILL
    c.font = INPUT_FONT
    c.border = ALL_BORDER
    if fmt:
        c.number_format = fmt


def _label(ws, row: int, col: int, text: str, bold: bool = False) -> None:
    c = ws.cell(row=row, column=col, value=text)
    if bold:
        c.font = BOLD_FONT


def _formula_maybe(ws, row: int, col: int, formula: str, fmt: str | None) -> None:
    c = ws.cell(row=row, column=col, value=formula)
    c.border = ALL_BORDER
    if fmt:
        c.number_format = fmt


@function_tool(strict_mode=False)
def build_precedent_transactions_workbook(
    deal_context_name: str,
    transactions_json: str,
    output_file: str | None = None,
) -> str:
    """Build an Excel workbook spreading precedent M&A transactions (manual inputs).

    ``transactions_json`` must be a JSON array of objects. Recognised fields (aliases
    accepted): announce_date/date; target; acquirer/buyer; buyer_type ("strategic"
    / "financial"); enterprise_value/ev_usd; ltm_ebitda/ebitda_usd; ev_ebitda_multiple
    / multiple; control_premium_pct/premium_pct; notes.

    Computes median EV/EBITDA on Cover via formulas referencing the Transactions sheet.
    """
    try:
        rows_in = json.loads(transactions_json)
    except json.JSONDecodeError as e:
        return json.dumps({"status": "error", "message": f"Invalid JSON: {e}"})

    if not isinstance(rows_in, list):
        return json.dumps({"status": "error", "message": "transactions_json must be a JSON array"})

    path = _out_path(output_file, "Precedent_Transactions.xlsx")
    wb = Workbook()
    wb.remove(wb.active)

    tx = wb.create_sheet("Transactions")
    hdr = [
        "Announce date",
        "Target",
        "Acquirer",
        "Buyer type",
        "Enterprise value ($)",
        "LTM EBITDA ($)",
        "EV / EBITDA (x)",
        "Control premium %",
        "Notes",
    ]
    for j, h in enumerate(hdr, start=1):
        c = tx.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = ALL_BORDER

    norm_rows: list[dict[str, Any]] = []
    for raw in rows_in:
        if not isinstance(raw, dict):
            continue
        d = raw.get("announce_date") or raw.get("date") or ""
        tgt = raw.get("target") or raw.get("Target") or ""
        acq = raw.get("acquirer") or raw.get("buyer") or ""
        bt = str(raw.get("buyer_type") or raw.get("Buyer type") or "").lower()
        if "strategic" in bt:
            bdisp = "Strategic"
        elif "financial" in bt or "sponsor" in bt or "pe" in bt:
            bdisp = "Financial"
        elif bt:
            bdisp = bt.title()
        else:
            bdisp = ""
        ev = raw.get("enterprise_value")
        if ev is None:
            ev = raw.get("ev_usd")
        eb = raw.get("ltm_ebitda")
        if eb is None:
            eb = raw.get("ebitda_usd")
        mult = raw.get("ev_ebitda_multiple")
        if mult is None:
            mult = raw.get("multiple")
        prem = raw.get("control_premium_pct")
        if prem is None:
            prem = raw.get("premium_pct")
        notes = raw.get("notes") or ""
        try:
            ev_f = float(ev) if ev is not None and ev != "" else None
        except (TypeError, ValueError):
            ev_f = None
        try:
            eb_f = float(eb) if eb is not None and eb != "" else None
        except (TypeError, ValueError):
            eb_f = None
        try:
            m_f = float(mult) if mult is not None and mult != "" else None
        except (TypeError, ValueError):
            m_f = None
        try:
            p_f = float(prem) if prem is not None and prem != "" else None
        except (TypeError, ValueError):
            p_f = None
        norm_rows.append(
            {
                "d": d,
                "tgt": tgt,
                "acq": acq,
                "bdisp": bdisp,
                "ev": ev_f,
                "eb": eb_f,
                "mult": m_f,
                "prem": p_f,
                "notes": notes,
            }
        )

    if not norm_rows:
        norm_rows.append(
            {
                "d": "",
                "tgt": "",
                "acq": "",
                "bdisp": "",
                "ev": None,
                "eb": None,
                "mult": None,
                "prem": None,
                "notes": "Add transaction rows via agent or paste JSON and rebuild",
            }
        )
    r = 2
    for nr in norm_rows:
        _input_cell(tx, r, 1, nr["d"])
        _input_cell(tx, r, 2, nr["tgt"])
        _input_cell(tx, r, 3, nr["acq"])
        _input_cell(tx, r, 4, nr["bdisp"])
        _input_cell(tx, r, 5, nr["ev"], FMT_USD if nr["ev"] else None)
        _input_cell(tx, r, 6, nr["eb"], FMT_USD if nr["eb"] else None)
        if nr["mult"] is not None:
            _input_cell(tx, r, 7, nr["mult"], FMT_MULT)
        else:
            _formula_maybe(
                tx,
                r,
                7,
                f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r}),F{r}<>0),E{r}/F{r},NA())',
                FMT_MULT,
            )
        _input_cell(tx, r, 8, nr["prem"], FMT_PCT if nr["prem"] is not None else None)
        _input_cell(tx, r, 9, nr["notes"])
        r += 1

    last_row = max(2, r - 1)
    autosize_columns(tx)

    summ = wb.create_sheet("Summary")
    summ["A1"] = f"Precedent transactions — {deal_context_name}"
    summ["A1"].font = Font(size=14, bold=True, color="1F4E78")
    summ.merge_cells("A1:F1")
    summ["A3"] = "Deal count"
    summ["B3"] = f"=COUNTA(Transactions!B2:B{last_row})"
    summ["A4"] = "Median EV/EBITDA (live — ignores errors)"
    hi = max(last_row, 2)
    summ["B4"] = f"=AGGREGATE(12,6,Transactions!G2:G{hi})"
    summ["B4"].number_format = FMT_MULT

    mults = [x["mult"] for x in norm_rows if x["mult"] is not None]
    if not mults:
        implied = []
        for x in norm_rows:
            if x["ev"] and x["eb"] and x["eb"] > 0:
                implied.append(x["ev"] / x["eb"])
        mults = implied

    cover = wb.create_sheet("Cover", 0)
    cover["A1"] = deal_context_name
    cover["A1"].font = Font(size=18, bold=True, color="1F4E78")
    cover["A2"] = "Precedent transaction spread"
    cover["A2"].font = Font(italic=True, color="666666")
    cover["A4"] = "Median EV/EBITDA (from Transactions tab)"
    cover["B4"] = "=Summary!B4"
    cover["B4"].number_format = FMT_MULT
    cover["A5"] = "Python cross-check median (JSON headline)"
    if mults:
        cover["B5"] = float(statistics.median(mults))
        cover["B5"].number_format = FMT_MULT
    cover["A7"] = "Instructions"
    cover["B7"] = (
        "Paste or edit yellow cells. Buyer type helps segment strategic vs financial buyers "
        "in narrative comps — no automatic split median here (extend if needed)."
    )

    wb.save(path)

    return json.dumps(
        {
            "status": "ok",
            "path": str(path),
            "deal_context_name": deal_context_name,
            "transactions_written": len(norm_rows),
            "median_ev_ebitda_approx": float(statistics.median(mults)) if mults else None,
        },
        indent=2,
    )


@function_tool(strict_mode=False)
def build_diligence_tracker_workbook(
    deal_name: str,
    output_file: str | None = None,
    extra_workstreams_json: str | None = None,
) -> str:
    """Create a diligence workstream tracker (Excel) with standard PE columns.

    ``extra_workstreams_json``: optional JSON array of strings appended after defaults.
    """
    defaults = [
        "Commercial / Market",
        "Financial / QoE",
        "Tax",
        "Legal / Corporate",
        "HR / Compensation",
        "IT / Cybersecurity",
        "Environmental / ESG",
        "Insurance",
        "Operational",
        "Anti-corruption / Sanctions",
        "Debt financing / Consents",
    ]
    extra: list[str] = []
    if extra_workstreams_json:
        try:
            parsed = json.loads(extra_workstreams_json)
            if isinstance(parsed, list):
                extra = [str(x) for x in parsed]
        except json.JSONDecodeError:
            pass

    ws_rows = defaults + extra
    path = _out_path(output_file, f"{deal_name[:40]}_Diligence_Tracker.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Diligence"
    ws["A1"] = f"Diligence tracker — {deal_name}"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")

    headers = [
        "Workstream",
        "Lead",
        "Target date",
        "Status",
        "Priority",
        "Key questions",
        "Findings summary",
        "Red flag?",
        "Mitigants",
    ]
    hr = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=j, value=h).fill = HEADER_FILL
        ws.cell(row=hr, column=j).font = HEADER_FONT
        ws.cell(row=hr, column=j).border = ALL_BORDER

    r = hr + 1
    for name in ws_rows:
        _label(ws, r, 1, name, bold=True)
        for c in range(2, 10):
            _input_cell(ws, r, c, "")
        r += 1

    autosize_columns(ws, min_width=14)
    wb.save(path)

    return json.dumps({"status": "ok", "path": str(path), "workstreams": len(ws_rows)}, indent=2)


@function_tool(strict_mode=False)
def compare_financing_facilities_workbook(
    facility_labels_json: str,
    terms_matrix_json: str,
    output_file: str | None = None,
) -> str:
    """Side-by-side financing comparison workbook.

    ``facility_labels_json``: JSON array of strings, e.g. ``["TLB Lead arr.", "Unitranche"]``.
    ``terms_matrix_json``: JSON object mapping row label → array of values (one per facility),
    e.g. ``{"Maturity (years)":[7,6],"Margin (bps)":[450,475],"Floor":[0,0],"OID %":[2,1.5]}``.
    """
    try:
        labels = json.loads(facility_labels_json)
        matrix = json.loads(terms_matrix_json)
    except json.JSONDecodeError as e:
        return json.dumps({"status": "error", "message": str(e)})

    if not isinstance(labels, list) or not labels:
        return json.dumps({"status": "error", "message": "facility_labels_json must be a non-empty array"})
    if not isinstance(matrix, dict):
        return json.dumps({"status": "error", "message": "terms_matrix_json must be an object"})

    labels = [str(x) for x in labels]
    nfac = len(labels)

    path = _out_path(output_file, "Financing_Term_Comparison.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Compare"
    ws["A1"] = "Financing term comparison"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")

    ws.cell(row=3, column=1, value="Term").fill = HEADER_FILL
    ws.cell(row=3, column=1).font = HEADER_FONT
    for j, lab in enumerate(labels, start=2):
        c = ws.cell(row=3, column=j, value=lab)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    r = 4
    for row_label in matrix.keys():
        vals = matrix[row_label]
        if not isinstance(vals, list):
            vals = [vals]
        while len(vals) < nfac:
            vals.append("")
        vals = vals[:nfac]
        ws.cell(row=r, column=1, value=str(row_label)).font = BOLD_FONT
        for j in range(nfac):
            _input_cell(ws, r, j + 2, vals[j])
        r += 1

    autosize_columns(ws, min_width=16)
    wb.save(path)

    return json.dumps({"status": "ok", "path": str(path), "facilities": nfac, "term_rows": r - 4}, indent=2)


@function_tool(strict_mode=False)
def build_portfolio_monitoring_workbook(
    portfolio_company_name: str,
    output_file: str | None = None,
    kpi_labels_json: str | None = None,
) -> str:
    """Board-pack style monitoring template: KPI grid vs budget / prior year.

    ``kpi_labels_json``: optional JSON array of KPI names; defaults to revenue, EBITDA,
    margin %, cash, net debt, WC days.
    """
    default_kpis = [
        "Revenue",
        "EBITDA",
        "EBITDA margin %",
        "Cash",
        "Net debt",
        "NWC days",
    ]
    kpis = default_kpis
    if kpi_labels_json:
        try:
            p = json.loads(kpi_labels_json)
            if isinstance(p, list) and p:
                kpis = [str(x) for x in p]
        except json.JSONDecodeError:
            pass

    path = _out_path(output_file, f"{portfolio_company_name[:30]}_Portfolio_Monitor.xlsx")
    wb = Workbook()

    cov = wb.active
    cov.title = "Cover"
    cov["A1"] = portfolio_company_name
    cov["A1"].font = Font(size=16, bold=True, color="1F4E78")
    cov["A2"] = "Portfolio monitoring template — populate monthly / quarterly"
    cov["A4"] = "Reporting period"
    _input_cell(cov, 4, 2, "")
    cov["A5"] = "Prepared by"
    _input_cell(cov, 5, 2, "")
    cov["A6"] = "vs budget variance commentary"
    _input_cell(cov, 6, 2, "")

    kpi = wb.create_sheet("KPIs")
    periods = ["Prior year", "Budget", "Actual", "Variance vs budget", "Variance vs PY"]
    kpi["A1"] = "KPI"
    kpi["A1"].font = BOLD_FONT
    for j, p in enumerate(periods, start=2):
        c = kpi.cell(row=1, column=j, value=p)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    for i, name in enumerate(kpis, start=2):
        kpi.cell(row=i, column=1, value=name).font = BOLD_FONT
        for j in range(2, 2 + len(periods)):
            _input_cell(kpi, i, j, "")

    act = wb.create_sheet("Actions")
    act["A1"] = "Board / management actions"
    act["A1"].font = Font(size=12, bold=True)
    hdr = ["Action", "Owner", "Due", "Status"]
    for j, h in enumerate(hdr, start=1):
        act.cell(row=3, column=j, value=h).fill = HEADER_FILL
        act.cell(row=3, column=j).font = HEADER_FONT
    for rr in range(4, 14):
        for cc in range(1, 5):
            _input_cell(act, rr, cc, "")

    autosize_columns(kpi)
    autosize_columns(act)
    autosize_columns(cov)
    wb.save(path)

    return json.dumps({"status": "ok", "path": str(path), "kpi_rows": len(kpis)}, indent=2)


@function_tool(strict_mode=False)
def build_covenant_review_workbook(
    deal_name: str,
    output_file: str | None = None,
    instrument_name: str = "Senior secured facilities",
) -> str:
    """Covenant review checklist — typical credit-agreement themes (manual excerpts).

    Yellow cells for thresholds, headroom, and compliance sign-off.
    """
    path = _out_path(output_file, f"{deal_name[:35]}_Covenant_Review.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Covenants"

    ws["A1"] = f"Covenant review — {deal_name}"
    ws["A2"] = instrument_name
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")

    headers = [
        "Covenant theme",
        "Clause / covenant name",
        "Covenant text excerpt",
        "Formula / threshold",
        "Step-down schedule?",
        "Last tested level",
        "Headroom / cushion",
        "Compliant? (Y/N)",
        "Notes / mitigation",
    ]
    hr = 4
    style_header_row(ws, hr, len(headers))
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=j, value=h)

    themes = [
        ("Leverage", "Total Net Leverage Ratio"),
        ("Leverage", "Senior Secured Leverage"),
        ("Coverage", "Interest Coverage Ratio"),
        ("Coverage", "Fixed Charge Coverage"),
        ("Basket", "Restricted Payments"),
        ("Basket", "Indebtedness incurrence"),
        ("Basket", "Liens"),
        ("Operational", "Asset sale proceeds sweep"),
        ("Operational", "M&A / JVs limitations"),
        ("Reporting", "Financial reporting deadlines"),
        ("Reporting", "Compliance certificate"),
        ("Events", "MAC / cross-default sensitivity"),
    ]

    r = hr + 1
    for theme, cname in themes:
        ws.cell(row=r, column=1, value=theme).font = BOLD_FONT
        ws.cell(row=r, column=2, value=cname).font = BOLD_FONT
        for c in range(3, 10):
            _input_cell(ws, r, c, "")
        r += 1

    autosize_columns(ws, min_width=18)
    wb.save(path)

    return json.dumps({"status": "ok", "path": str(path), "covenant_rows": len(themes)}, indent=2)
